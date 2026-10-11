"""Build six editable Word documents from the maintained bilingual Markdown.

Requires python-docx and Pillow. Run: python 04_manuals/build_word_documents.py
Word edits do not update the Markdown sources: reconcile them before rebuilding.
The validation JSON records source/output hashes and structural checks; visual
pagination must also be checked in a Word-compatible renderer before release.
"""
from pathlib import Path
import hashlib
import json
import re
import unicodedata
import argparse

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

R = Path(__file__).resolve().parents[1]
WIDTH = 170
BUILD_DATE = '2026-10-08'
RECORDS = []


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def heading_text(text):
    # Preserve technical punctuation in body text; headings are plain words.
    text = re.sub(r'[*`]', '', text)
    return re.sub(r'\s+', ' ', ''.join(' ' if unicodedata.category(c).startswith('P') else c for c in text)).strip()


def set_font(run, size=None, bold=None):
    run.font.name = 'Calibri'
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft JhengHei')
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_font(style, size, bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        rf.attrib.pop(qn(k), None)


def document(title, lang):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
    sec.left_margin, sec.right_margin = Mm(20), Mm(20)
    sec.header_distance, sec.footer_distance = Mm(8), Mm(9)
    normal = doc.styles['Normal']
    style_font(normal, 10.5 if lang == 'ZH' else 11)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 23), ('Subtitle', 11), ('Heading 1', 15), ('Heading 2', 12), ('Heading 3', 11)]:
        st = doc.styles[name]
        style_font(st, size, name != 'Subtitle')
        st.paragraph_format.space_before = Pt(12 if name.startswith('Heading') else 0)
        st.paragraph_format.space_after = Pt(7)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.keep_together = True
        st.paragraph_format.line_spacing = 1.12
        borders = st.element.get_or_add_pPr().find(qn('w:pBdr'))
        if borders is not None:
            borders.getparent().remove(borders)
    for name in ['Header', 'Footer', 'Caption']:
        style_font(doc.styles[name], 8.5)
    doc.core_properties.title = title
    doc.core_properties.subject = 'AeroSense engineering prototype editable documentation'
    doc.core_properties.author = 'AeroSense project team'
    doc.core_properties.keywords = 'iGEM AeroSense editable manual'
    p = sec.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run('AeroSense  •  ' + ('工程原型' if lang == 'ZH' else 'Engineering prototype')), 8)
    f = sec.footer.paragraphs[0]
    f.paragraph_format.tab_stops.add_tab_stop(Mm(WIDTH), WD_ALIGN_PARAGRAPH.RIGHT)
    set_font(f.add_run(BUILD_DATE + '   |   ' + ('可編輯文件' if lang == 'ZH' else 'Editable document') + '\t'), 8)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    f._p.append(fld)
    doc.add_paragraph(heading_text(title), 'Title')
    return doc


def add_hyperlink(p, text, target):
    link = OxmlElement('w:hyperlink')
    rel = p.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '125D75'); rp.append(c)
    fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'), 'Calibri'); fonts.set(qn('w:eastAsia'), 'Microsoft JhengHei'); rp.append(fonts)
    r.append(rp)
    t = OxmlElement('w:t'); t.text = text; r.append(t)
    link.append(r); p._p.append(link)


def inline(p, text, size=None):
    text = re.sub(r'□[ \t]+', '□\u00a0', text)
    pattern = r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))'
    for token in re.split(pattern, text):
        if not token:
            continue
        m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if m:
            if m[2].startswith(('http://', 'https://')):
                add_hyperlink(p, m[1], m[2])
            else:
                run = p.add_run(m[1] + ' (' + m[2] + ')'); set_font(run, size)
        elif token.startswith('**'):
            set_font(p.add_run(token[2:-2]), size, True)
        elif token.startswith('`'):
            run = p.add_run(token[1:-1]); set_font(run, size or 9.5)
            run.font.name = 'Consolas'
        else:
            set_font(p.add_run(token), size)


def para(doc, text, size=None):
    p = doc.add_paragraph()
    inline(p, text, size)
    return p


def table(doc, rows, widths=None, size=9.5, minimum_height=None):
    n = len(rows[0])
    widths = widths or ({2: [57, 113], 3: [41, 80, 49], 4: [30, 49, 49, 42]}.get(n, [WIDTH / n] * n))
    t = doc.add_table(rows=0, cols=n)
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Mm(width)
    borders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = OxmlElement('w:' + name); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
    t._tbl.tblPr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side, val in [('top', 95), ('bottom', 95), ('left', 100), ('right', 100)]:
        e = OxmlElement('w:' + side); e.set(qn('w:w'), str(val)); e.set(qn('w:type'), 'dxa'); margins.append(e)
    t._tbl.tblPr.append(margins)
    for index, values in enumerate(rows):
        row = t.add_row()
        trp = row._tr.get_or_add_trPr()
        trp.append(OxmlElement('w:cantSplit'))
        if index == 0:
            e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trp.append(e)
        if minimum_height and index:
            row.height = Mm(minimum_height); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Mm(widths[j]); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            inline(p, value, size)
            if n >= 6 and j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if index == 0:
                p.paragraph_format.keep_with_next = True
                for run in p.runs:
                    run.bold = True
                shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'E8EEF1'); cell._tc.get_or_add_tcPr().append(shd)
    spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(3); spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.35
    return t


def list_paragraph(doc, value, ordered=False):
    # Retain the explicit step numbering from the maintained source.
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Mm(6)
    p.paragraph_format.first_line_indent = Mm(-6)
    p.paragraph_format.keep_together = True
    inline(p, value if ordered else '• ' + value)
    return p


def markdown(doc, path, assembly=False, first_title=True):
    lines = path.read_text(encoding='utf-8-sig').splitlines()
    i = 0; images = 0; paragraph_count = 0
    while i < len(lines):
        s = lines[i].strip(); i += 1
        if not s:
            continue
        if s.startswith('```'):
            while i < len(lines) and not lines[i].startswith('```'):
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.08
                p.paragraph_format.keep_together = True
                run = p.add_run(lines[i]); set_font(run, 8.5); run.font.name = 'Consolas'
                i += 1
            i += 1
            continue
        if s.startswith('#'):
            level = len(s) - len(s.lstrip('#')); text = heading_text(s[level:].strip())
            if level == 1 and first_title:
                first_title = False
                continue
            p = doc.add_paragraph(text, 'Heading 1' if level <= 2 else 'Heading 2')
            if False:  # Continuous task layout avoids artificial blank pages.
                p.paragraph_format.page_break_before = True
            continue
        m = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', s)
        if m:
            imagepath = (path.parent / m[2]).resolve()
            w, h = Image.open(imagepath).size
            scale = min(165 / w, 74 / h)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(imagepath), width=Mm(w * scale), height=Mm(h * scale))
            cap = doc.add_paragraph(m[1], 'Caption'); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.keep_with_next = True
            cap.paragraph_format.space_after = Pt(4)
            detail = doc.add_paragraph(style='Caption')
            detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
            label = '開啟原尺寸圖查看細部標註' if '_ZH' in path.stem else 'Open full-size drawing for detailed labels'
            add_hyperlink(detail, label + '  ' + imagepath.name, m[2])
            detail.paragraph_format.space_after = Pt(4)
            images += 1
            continue
        if s.startswith('|') and i < len(lines) and re.match(r'^\|[\s:|\-]+\|?$', lines[i].strip()):
            rows = [[x.strip() for x in s.strip('|').split('|')]]; i += 1
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([x.strip() for x in lines[i].strip().strip('|').split('|')]); i += 1
            table(doc, rows)
            continue
        if re.match(r'^\d+\.\s', s):
            list_paragraph(doc, s, True)
        elif s.startswith('- '):
            list_paragraph(doc, s[2:])
        elif s in ['---', '***']:
            continue
        else:
            para(doc, s.lstrip('> '))
        paragraph_count += 1
    return {'source': path.relative_to(R).as_posix(), 'SHA256': digest(path), 'images': images, 'body_paragraphs': paragraph_count}


def save(doc, target, sources, kind, lang):
    doc.save(target)
    reloaded = Document(target)
    assert reloaded.paragraphs[0].style.name == 'Title'
    assert not any(x.tag == qn('w:drawing') for x in reloaded.paragraphs[0]._p.iter())
    assert len(reloaded.paragraphs) > 15
    assert len(reloaded.inline_shapes) == sum(x.get('images', 0) for x in sources)
    RECORDS.append({'file': target.relative_to(R).as_posix(), 'kind': kind, 'language': lang,
                    'sources': sources, 'SHA256': digest(target), 'paragraphs': len(reloaded.paragraphs),
                    'tables': len(reloaded.tables), 'embedded_figures': len(reloaded.inline_shapes),
                    'editable_native_text_and_tables': True, 'protected': False})


def split_feedback(path):
    text = path.read_text(encoding='utf8')
    blocks = re.split(r'^## (.+)$', text, flags=re.M)
    intro = [s for s in blocks[0].splitlines() if s.strip() and not s.startswith('#')]
    sections = [(blocks[i], [s.strip() for s in blocks[i + 1].splitlines() if s.strip()]) for i in range(1, len(blocks), 2)]
    return intro, sections


def writing_space(doc, mm=20):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Mm(mm)
    return p


def questionnaire(lang):
    path = R / '07_validation' / f'User_Feedback_{lang}.md'
    title = 'AeroSense 組裝與使用回饋表' if lang == 'ZH' else 'AeroSense assembly and use feedback'
    doc = document(title, lang)
    if lang == 'ZH':
        # Fit the consent and eight ratings on one page without shrinking type.
        normal = doc.styles['Normal'].paragraph_format
        normal.space_after = Pt(4)
        normal.line_spacing = 1.08
        snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0')
        doc.styles['Normal'].element.get_or_add_pPr().append(snap)
        for name in ['Heading 1', 'Heading 2']:
            doc.styles[name].paragraph_format.space_before = Pt(8)
            doc.styles[name].paragraph_format.space_after = Pt(5)
    intro, sections = split_feedback(path)
    for s in intro:
        # Two lines avoid cramping the four metadata fields.
        if s.startswith(('參與者代碼', 'Participant code')):
            if lang == 'ZH':
                para(doc, '參與者代碼：________________　日期：________________')
                para(doc, '手冊版本：__________________　觀察者代碼：____________')
            else:
                para(doc, 'Participant code: ________________   Date: ________________')
                para(doc, 'Manual version: ________________   Observer code: ____________')
        else:
            para(doc, s, 10 if lang == 'EN' else 10.5)
    doc.add_paragraph(heading_text(sections[0][0]), 'Heading 1')
    for s in sections[0][1]:
        para(doc, s, 10 if lang == 'EN' else 10.5)
    ratings = sections[2][1]
    doc.add_paragraph(heading_text(sections[2][0]), 'Heading 1')
    para(doc, ratings[0], 9.5)
    ratingrows = [[x.strip() for x in s.strip('|').split('|')] for s in ratings if s.startswith('|') and not re.match(r'^\|[\s:|\-]+\|?$', s)]
    table(doc, ratingrows, [105, 10, 10, 10, 10, 10, 15], size=9.2, minimum_height=8)

    doc.add_page_break()
    doc.add_paragraph(heading_text(sections[1][0]), 'Heading 1')
    tasks = sections[1][1]
    para(doc, tasks[0], 10)
    taskrows = [s for s in tasks if s.startswith('| ')][1:]
    for s in taskrows:
        title = s.strip('|').split('|')[0].strip()
        p = para(doc, title, 10.5)
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            run.bold = True
        if lang == 'ZH':
            p = para(doc, '時間：______ 秒　提示：______ 次　□ 獨立完成　□ 協助後完成　□ 未完成', 9.5)
            p.paragraph_format.keep_with_next = True
            para(doc, '卡住或做錯的步驟：________________________________________________', 9.5)
        else:
            p = para(doc, 'Time: ______ s   Prompts: ____   □ Independent   □ Assisted   □ Incomplete', 9.5)
            p.paragraph_format.keep_with_next = True
            para(doc, 'Where stuck / error: __________________________________________________', 9.5)
        writing_space(doc, 4)
    for s in tasks[1:]:
        if not s.startswith('|'):
            para(doc, s, 10)

    doc.add_page_break()
    doc.add_paragraph('請用自己的話回饋' if lang == 'ZH' else 'Feedback in your own words', 'Heading 1')
    para(doc, '請寫具體步驟或例子。未操作的項目可留空。' if lang == 'ZH' else 'Give a specific step or example. Leave unperformed items blank.')
    open_questions = [s for s in ratings[1:] if not s.startswith('|')]
    for i, s in enumerate(open_questions, 1):
        label = re.sub(r'_+', '', s).strip()
        p = para(doc, str(i) + '. ' + label)
        p.paragraph_format.keep_with_next = True
        writing_space(doc, 29)

    doc.add_page_break()
    doc.add_paragraph(heading_text(sections[3][0]), 'Heading 1')
    for i, s in enumerate(sections[3][1]):
        if i == 0 and lang == 'EN':
            s = s.replace('Severity:', '\nSeverity:').replace(' □ Inconvenient.', '\n□ Inconvenient.')
        p = para(doc, s)
        if i in [1, 2, 3, 5]:
            p.paragraph_format.keep_with_next = True
            writing_space(doc, 21)
    save(doc, path.with_suffix('.docx'), [{'source': path.relative_to(R).as_posix(), 'SHA256': digest(path), 'images': 0}], 'feedback', lang)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--kind', nargs='+', choices=['assembly','operation','feedback'], default=['assembly','operation','feedback'])
    ap.add_argument('--language', nargs='+', choices=['ZH','EN'], default=['ZH','EN'])
    args = ap.parse_args()
    validation_path = R / '04_manuals/word_document_validation.json'
    previous = json.loads(validation_path.read_text(encoding='utf8'))['documents'] if validation_path.exists() else []
    for lang in ['ZH', 'EN']:
        if lang not in args.language: continue
        src = R / '04_manuals' / f'Assembly_Manual_{lang}.md'
        title = 'AeroSense 圖解組裝手冊' if lang == 'ZH' else 'AeroSense illustrated assembly manual'
        if 'assembly' in args.kind:
            doc = document(title, lang)
            records = [markdown(doc, src, assembly=True)]
            save(doc, src.with_suffix('.docx'), records, 'assembly', lang)

        src = R / '03_software/guides' / ('操作與設定_ZH.md' if lang == 'ZH' else 'Operation_and_Setup_EN.md')
        correction = src.parent / f'Measured_Correction_{lang}.md'
        title = 'AeroSense 操作與設定手冊' if lang == 'ZH' else 'AeroSense operation and setup manual'
        if 'operation' in args.kind:
            doc = document(title, lang)
            records = [markdown(doc, src)]
            doc.add_page_break()
            records.append(markdown(doc, correction, first_title=False))
            save(doc, R / '04_manuals' / f'Operation_and_Setup_{lang}.docx', records, 'operation', lang)
        if 'feedback' in args.kind: questionnaire(lang)
    generated = {x['file'] for x in RECORDS}
    RECORDS.extend(x for x in previous if x['file'] not in generated)
    report = {'date': BUILD_DATE, 'count': len(RECORDS), 'documents': RECORDS,
              'editable': True, 'automatic_word_to_markdown_sync': False,
              'visual_QA': 'See word_render_validation.json for the separately checked rendered pages.'}
    (R / '04_manuals/word_document_validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf8')
    print(json.dumps([{'file': x['file'], 'figures': x['embedded_figures'], 'tables': x['tables']} for x in RECORDS], ensure_ascii=False))


if __name__ == '__main__':
    main()
